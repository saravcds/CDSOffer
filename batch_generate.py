"""
First-cycle batch: generate offer letters for every employee in the CTC details sheet.

Usage:
  python batch_generate.py "Emp Details - Offers 2.xlsx" --date 2026-10-26 --out "Offers_2026_Cycle1"
"""
import argparse
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from offer_engine import (ROW_ORDER, Register, docx_to_pdf_many, fill_letter, inr, make_doc_id,
                          read_employees_xlsx, register_entry, safe_filename)

HERE = Path(__file__).parent


def write_review(emps, ids, path: Path, letter_date: date):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Review"
    cols = ["Doc ID", "Employee Name", "Dear ...", "Designation", "Address (as printed)", "Annual CTC"]
    cols += [f"{label} (Yearly)" for _, label in ROW_ORDER]
    cols += [f"{label} (Monthly)" for _, label in ROW_ORDER]
    cols += ["CTC ties?", "Check / Flags", "Source row"]
    ws.append(cols)
    head = PatternFill("solid", fgColor="1F3864")
    flag = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin", color="BFBFBF")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for e, did in zip(emps, ids):
        if e.ctc is None:
            ws.append([did, e.name, e.salutation, e.designation, "\n".join(e.address_lines), None]
                      + [None] * 18 + ["", "; ".join(e.flags), e.source_row])
        else:
            b = e.breakup
            y, m = b["yearly"], b["monthly"]
            tie = (y["total_a"] + y["total_b"] == y["ctc"])
            ws.append([did, e.name, e.salutation, e.designation, "\n".join(e.address_lines), e.ctc]
                      + [y[k] for k, _ in ROW_ORDER] + [m[k] for k, _ in ROW_ORDER]
                      + ["Yes" if tie else "NO", "; ".join(e.flags), e.source_row])
        r = ws.max_row
        for c in ws[r]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if isinstance(c.value, (int, float)) and c.column >= 6 and c.column <= 24:
                c.number_format = "#,##,##0"
        if e.flags:
            for c in ws[r]:
                c.fill = flag
    widths = [24, 28, 22, 30, 42, 13] + [13] * 18 + [9, 40, 8]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 45

    # Formula logic sheet, so HR can see exactly how numbers were made
    lg = wb.create_sheet("Formula used")
    for row in [
        ["Component", "Formula (same as 'Salary Structure' sheet)"],
        ["Basic Salary", "CTC x 50%"],
        ["House Rent Allowance", "Basic x 40%"],
        ["PF Employer Contribution", "21,600 per year (fixed)"],
        ["Gratuity", "Basic x 15 / 26 / 12"],
        ["Professional Tax", "2,400 per year (fixed)"],
        ["Additional Allowance", "CTC - (Basic + HRA + PF + Gratuity + PT)  -> balancing figure"],
        ["Monthly", "Yearly / 12, rounded to whole rupees per row"],
        ["Letter / joining / w.e.f. date", letter_date.strftime("%d %B %Y")],
        ["Generated on", datetime.now().strftime("%d %b %Y %H:%M")],
    ]:
        lg.append(row)
    lg["A1"].font = lg["B1"].font = Font(bold=True)
    lg.column_dimensions["A"].width = 30
    lg.column_dimensions["B"].width = 70
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--sheet", default="CTC details")
    ap.add_argument("--date", default="2026-10-26", help="letter / joining date YYYY-MM-DD")
    ap.add_argument("--out", default="Offers_2026_Cycle1")
    ap.add_argument("--template", default=str(HERE / "templates" / "CDS Appt Letter.docx"))
    ap.add_argument("--register", default=str(HERE / "data" / "offer_register.xlsx"))
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()

    letter_date = datetime.strptime(a.date, "%Y-%m-%d").date()
    year = letter_date.year
    out = Path(a.out)
    (out / "Word").mkdir(parents=True, exist_ok=True)

    emps = read_employees_xlsx(a.excel, a.sheet)
    reg = Register(a.register)
    seq0 = reg.next_seq(year)
    ids = [make_doc_id(year, seq0 + i) for i in range(len(emps))]

    write_review(emps, ids, out / f"Offer Review - {letter_date:%d %b %Y}.xlsx", letter_date)

    docx_files, entries = [], []
    for i, (e, did) in enumerate(zip(emps, ids)):
        if e.ctc is None:
            print(f"SKIP {e.name}: {e.flags}")
            continue
        fname = safe_filename(did, e.name)
        p = out / "Word" / f"{fname}.docx"
        p.write_bytes(fill_letter(a.template, e, did, letter_date))
        docx_files.append(p)
        entries.append(register_entry(e, did, year, seq0 + i, letter_date, letter_date, fname,
                                      status="Issued - check flags" if e.flags else "Issued"))
    if not a.no_pdf:
        docx_to_pdf_many(docx_files, out / "PDF")
    reg.add(entries)
    print(f"Done: {len(docx_files)} letters, IDs {ids[0]} to {ids[-1]}")


if __name__ == "__main__":
    main()
