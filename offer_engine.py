"""
CDS Asia offer letter engine.

Used by both the batch script (first cycle) and the Streamlit HR app.
- Reads employee rows (Employee Name, Annual CTC, Designation, Address)
- Calculates the salary breakup using the Salary Structure formulas
- Fills the appointment letter template (Word) and optionally converts to PDF
- Keeps an offer register so document IDs never repeat
"""
from __future__ import annotations

import copy
import io
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches

# ---------------------------------------------------------------------------
# Salary structure (mirrors the "Salary Structure" sheet in Emp Details - Offers 2)
#   Basic      = CTC * 50%
#   HRA        = Basic * 40%
#   PF (Er)    = 21,600 fixed
#   Gratuity   = Basic * 15 / 26 / 12
#   Prof. Tax  = 2,400 fixed
#   Additional = CTC - (Basic + HRA + PF + Gratuity + PT)   (balancing figure)
#   Monthly    = Yearly / 12
# ---------------------------------------------------------------------------
BASIC_PCT = 0.50
HRA_PCT_OF_BASIC = 0.40
PF_EMPLOYER_YEARLY = 21600
PROF_TAX_YEARLY = 2400

DOC_ID_PREFIX = "CDSA/HR/APPT"
DOC_ID_START = 101

ROW_ORDER = [
    ("basic", "Basic Salary"),
    ("hra", "House Rent Allowance"),
    ("additional", "Additional Allowance"),
    ("total_a", "Total (A)"),
    ("pf", "PF Employer Contribution"),
    ("gratuity", "Gratuity"),
    ("pt", "Professional Tax"),
    ("total_b", "Total (B)"),
    ("ctc", "Cost to Company (A + B)"),
]


def calc_breakup(ctc: float) -> dict:
    """Return yearly and monthly components (whole rupees). CTC always ties exactly."""
    ctc = round(float(ctc))
    basic = round(ctc * BASIC_PCT)
    hra = round(basic * HRA_PCT_OF_BASIC)
    pf = PF_EMPLOYER_YEARLY
    gratuity = round(basic * 15 / 26 / 12)
    pt = PROF_TAX_YEARLY
    additional = ctc - (basic + hra + pf + gratuity + pt)
    yearly = {
        "basic": basic,
        "hra": hra,
        "additional": additional,
        "total_a": basic + hra + additional,
        "pf": pf,
        "gratuity": gratuity,
        "pt": pt,
        "total_b": pf + gratuity + pt,
        "ctc": ctc,
    }
    # Monthly = Yearly / 12 per row, same as the Excel sheet
    monthly = {k: round(v / 12) for k, v in yearly.items()}
    return {"yearly": yearly, "monthly": monthly}


def inr(n: float) -> str:
    """Indian digit grouping: 1400000 -> 14,00,000"""
    n = int(round(n))
    neg = n < 0
    s = str(abs(n))
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ",".join(parts) + "," + tail
    return ("-" if neg else "") + s


# ---------------------------------------------------------------------------
# Input cleaning
# ---------------------------------------------------------------------------
def parse_ctc(value) -> float | None:
    """'14LPA', '15 LPA', '16LPA  ', '1CrPA', 1400000 -> rupees"""
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        v = float(value)
        return v * 100000 if v < 1000 else v  # 14 -> 14 LPA
    s = str(value).strip().upper().replace(",", "").replace(" ", "")
    m = re.match(r"^([\d.]+)(LPA|L|LAKHS?|LAC|CRPA|CR|CRORES?)?(PA)?$", s)
    if not m:
        return None
    num = float(m.group(1))
    unit = m.group(2) or ""
    if unit.startswith("CR"):
        return num * 10000000
    if unit.startswith("L"):
        return num * 100000
    return num * 100000 if num < 1000 else num


KEEP_UPPER = {"QA", "AI", "HR", "IT", "CRM", "SVP", "UX", "D365", "VP", "BA", "PG", "RWD", "SS", "HMDA", "RTO", "CBCID", "DT", "AP"}
SMALL = {"and", "&", "of", "the", "for"}
SPELL_FIX = {"techinal": "Technical"}


def smart_title(text: str) -> str:
    out = []
    all_caps = text.isupper()
    for i, w in enumerate(text.split()):
        low = w.lower()
        if low in SPELL_FIX:
            out.append(SPELL_FIX[low])
        elif w.upper() in KEEP_UPPER and (w.isupper() or len(w) <= 4):
            out.append(w.upper())
        elif w.islower() or (all_caps and len(w) > 2 and not any(c.isdigit() for c in w)):
            out.append(low if (i > 0 and low in SMALL) else w[:1].upper() + w[1:].lower())
        else:
            out.append(w)
    return " ".join(out)


def clean_name(name: str) -> str:
    return smart_title(re.sub(r"\s+", " ", str(name or "")).strip())


def clean_designation(d: str) -> str:
    d = re.sub(r"\s+", " ", str(d or "")).strip()
    d = re.sub(r"\s*-\s*", " - ", d)
    return smart_title(d)


def salutation_name(full_name: str) -> str:
    """First and last name (drops middle names / initials)."""
    parts = full_name.split()
    if len(parts) <= 2:
        return full_name
    return f"{parts[0]} {parts[-1]}"


STATES = {
    "andhra pradesh", "telangana", "karnataka", "tamil nadu", "west bengal", "maharashtra", "kerala",
    "odisha", "gujarat", "delhi", "uttar pradesh", "rajasthan", "madhya pradesh", "bihar", "assam",
    "punjab", "haryana", "uttarakhand", "jharkhand", "chhattisgarh", "goa", "telengana",
}
PIN_RE = re.compile(r"^\d{6}$")


def clean_address(addr: str) -> list[str]:
    """Returns cleaned address lines. Removes '-' placeholders and repeated lines,
    and joins city / state / country / PIN onto one line."""
    if not addr or not str(addr).strip():
        return []
    raw = [re.sub(r"\s+", " ", l).strip().strip(",").strip() for l in str(addr).replace("\r", "\n").split("\n")]
    lines: list[str] = []
    for l in raw:
        if not l or l in {"-", "--", "."}:
            continue
        low = l.lower()
        if any(low == x.lower() or (len(low) > 3 and low in x.lower()) for x in lines):
            continue  # duplicate of something already in the address
        lines.append(smart_title(l))

    # join trailing city/state/country/pin
    tail: list[str] = []
    while lines:
        low = lines[-1].lower()
        if PIN_RE.match(lines[-1]) or re.match(r"^\d{5,8}$", lines[-1]) or low == "india" or low in STATES:
            tail.insert(0, lines.pop())
        else:
            break
    if tail and lines and any(t.lower() in STATES for t in tail) and len(lines[-1]) <= 25 and len(lines) > 1:
        tail.insert(0, lines.pop())
    if tail:
        pin = tail[-1] if re.match(r"^\d{5,8}$", tail[-1]) else None
        words = tail[:-1] if pin else tail
        joined = ", ".join(words)
        joined = f"{joined} - {pin}" if pin and joined else (pin or joined)
        lines.append(joined)
    return lines


def address_flags(addr_lines: list[str], raw: str) -> list[str]:
    flags = []
    if not addr_lines:
        flags.append("Address missing")
        return flags
    pins = re.findall(r"(?<!\d)\d{5,8}(?!\d)", " ".join(addr_lines))
    six = [p for p in pins if len(p) == 6]
    if not six and not pins:
        flags.append("No PIN code in address")
    elif not six:
        flags.append(f"PIN looks wrong ({pins[-1]})")
    return flags


@dataclass
class Employee:
    name: str
    designation: str
    ctc: float | None
    address_lines: list[str]
    salutation: str = ""
    joining_date: date | None = None
    doc_id: str = ""
    flags: list[str] = field(default_factory=list)
    source_row: int | None = None

    @property
    def breakup(self) -> dict:
        return calc_breakup(self.ctc)


def employee_from_row(name, ctc_raw, designation, address, row_no=None, salutation=None, joining=None) -> Employee:
    nm = clean_name(name)
    ctc = parse_ctc(ctc_raw)
    des = clean_designation(designation)
    lines = clean_address(address)
    flags = []
    if ctc is None:
        flags.append(f"CTC not readable ('{ctc_raw}')")
    if not des:
        flags.append("Designation missing")
    flags += address_flags(lines, address)
    if ctc and ctc >= 10000000:
        flags.append("CTC is 1 Cr or more - please confirm")
    if ctc and re.search(r"junior|trainee|intern", des, re.I) and ctc > 1000000:
        flags.append("High CTC for a junior role - please confirm")
    if ctc and calc_breakup(ctc)["yearly"]["additional"] < 0:
        flags.append("Additional Allowance is negative - CTC too low for this structure")
    return Employee(
        name=nm,
        designation=des,
        ctc=ctc,
        address_lines=lines,
        salutation=(clean_name(salutation) if salutation else salutation_name(nm)),
        joining_date=joining,
        flags=flags,
        source_row=row_no,
    )


def read_employees_xlsx(path_or_buffer, sheet_name: str = "CTC details") -> list[Employee]:
    wb = openpyxl.load_workbook(path_or_buffer, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[-1]
    headers = [str(c.value or "").strip().lower() for c in ws[1]]

    def col(*names):
        for n in names:
            for i, h in enumerate(headers):
                if h == n or h.startswith(n):
                    return i
        return None

    ci = {
        "name": col("employee name", "name"),
        "ctc": col("annual ctc", "ctc"),
        "des": col("designation"),
        "addr": col("address"),
        "sal": col("salutation"),
        "doj": col("date of joining", "joining date", "doj"),
    }
    if ci["name"] is None or ci["ctc"] is None:
        raise ValueError("Sheet needs at least 'Employee Name' and 'Annual CTC' columns")
    out = []
    for r_i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        g = lambda k: (row[ci[k]] if ci[k] is not None and ci[k] < len(row) else None)
        if not g("name") or not str(g("name")).strip():
            continue
        doj = g("doj")
        if isinstance(doj, datetime):
            doj = doj.date()
        elif not isinstance(doj, date):
            doj = None
        out.append(employee_from_row(g("name"), g("ctc"), g("des"), g("addr"), r_i, g("sal"), doj))
    return out


# ---------------------------------------------------------------------------
# Dates
# ---------------------------------------------------------------------------
def ordinal(n: int) -> str:
    return "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")


def fmt_long(d: date) -> str:  # 26 October 2026
    return f"{d.day} {d.strftime('%B')} {d.year}"


def fmt_dash(d: date) -> str:  # 26-October-2026
    return f"{d.day}-{d.strftime('%B')}-{d.year}"


# ---------------------------------------------------------------------------
# Word template filling
# ---------------------------------------------------------------------------
def _set_runs(paragraph, texts: list[str]):
    """Keep the first run's formatting, put text into it, clear the others."""
    runs = paragraph.runs
    for i, r in enumerate(runs):
        r.text = texts[0] if i == 0 else ""
    for t in texts[1:]:
        new = copy.deepcopy(runs[0]._r)
        paragraph._p.append(new)
    for i, t in enumerate(texts):
        paragraph.runs[i].text = t


def _replace_in_runs(paragraph, old: str, new: str) -> bool:
    """Replace text that may be split across runs, keeping the formatting of the run where it starts."""
    full = "".join(r.text for r in paragraph.runs)
    idx = full.find(old)
    if idx < 0:
        return False
    pos = 0
    start_run = None
    for r in paragraph.runs:
        r_start, r_end = pos, pos + len(r.text)
        pos = r_end
        if r_end <= idx or r_start >= idx + len(old):
            continue
        s = max(idx, r_start) - r_start
        e = min(idx + len(old), r_end) - r_start
        if start_run is None:
            start_run = r
            r.text = r.text[:s] + new + r.text[e:]
        else:
            r.text = r.text[:s] + r.text[e:]
    return True


def _cell_set(cell, text: str, bold: bool | None = None):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
        if bold is not None:
            p.runs[0].bold = bold
    else:
        r = p.add_run(text)
        if bold is not None:
            r.bold = bold


def _add_tab_stop(paragraph, pos_inches: float):
    ts = paragraph.paragraph_format.tab_stops
    ts.clear_all()
    ts.add_tab_stop(Inches(pos_inches), WD_TAB_ALIGNMENT.LEFT)


def fill_letter(template_path: str | os.PathLike, emp: Employee, doc_id: str, letter_date: date,
                joining_date: date | None = None, signatory: str = "Khushboo Kaur",
                signatory_title: str = "Head – HR") -> bytes:
    joining_date = joining_date or emp.joining_date or letter_date
    doc = Document(str(template_path))
    b = emp.breakup

    # 1. Doc ID + date (first table at top)
    top = doc.tables[0]
    _set_runs(top.cell(0, 0).paragraphs[0], [doc_id])
    _set_runs(top.cell(0, 1).paragraphs[0], [fmt_long(letter_date)])

    paras = doc.paragraphs
    for p in paras:
        t = p.text
        if t.strip() == "Emp Name":
            _set_runs(p, [emp.name])
        elif t.strip() == "Address":
            # one run with line breaks, same formatting
            for i, r in enumerate(p.runs):
                r.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT  # template is justified; breaks would stretch words
            run = p.runs[0]
            for i, line in enumerate(emp.address_lines or [""]):
                if i:
                    run.add_break()
                run.add_text(line)
        elif t.strip().startswith("Dear "):
            _replace_in_runs(p, "xxxx", emp.salutation)
        elif "pleased to offer you employment" in t:
            _replace_in_runs(p, "Designation ", emp.designation + " ")
        elif "this appointment is effective" in t:
            m = re.search(r"\d{1,2}-[A-Za-z]+-\d{4}", t)
            if m:
                _replace_in_runs(p, m.group(0), fmt_dash(joining_date))
        elif t.strip().startswith("I Employee Name"):
            _replace_in_runs(p, "Employee Name", emp.name)
        elif "(w.e.f." in t:
            d = joining_date
            runs = p.runs
            base = runs[0]
            # rebuild: "       (w.e.f. 26" + sup "th" + " October 2026)"
            sup_run = next((r for r in runs if r.font.superscript), None)
            for r in runs:
                r.text = ""
            runs[0].text = f"       (w.e.f. {d.day}"
            if sup_run is not None:
                sup_run.text = ordinal(d.day)
                # put remaining text in the run after the superscript
                after = runs[runs.index(sup_run) + 1] if runs.index(sup_run) + 1 < len(runs) else None
                if after is not None:
                    after.text = f" {d.strftime('%B')} {d.year})"
                else:
                    sup_run.text += f" {d.strftime('%B')} {d.year})"
            else:
                runs[0].text += f"{ordinal(d.day)} {d.strftime('%B')} {d.year})"
        elif t.startswith("For") and "ACCEPT" in t:
            # "For CDS Asia Pvt. Ltd.<tab>ACCEPTED"
            for i, r in enumerate(p.runs):
                r.text = ""
            p.runs[0].text = "For CDS Asia Pvt. Ltd."
            p.runs[0].add_tab()
            p.runs[0].add_text("ACCEPTED")
            _add_tab_stop(p, 4.4)
        elif t.strip().startswith(signatory.split()[0]) and "HR" in t:
            # "Khushboo Kaur<tab>Employee Name" / "Head – HR"
            runs = p.runs
            # run 0 = signatory name (bold); later runs hold the line break + title (normal)
            run = runs[0]
            run.text = signatory
            run.add_tab()
            run.add_text(emp.name)
            rest = [r for r in runs[1:] if r.text != "\n"]
            if rest and signatory_title != "Head – HR":
                rest[0].text = signatory_title
                for r in rest[1:]:
                    r.text = ""
            _add_tab_stop(p, 4.4)

    # 2. Salary annexure table
    ann = doc.tables[1]
    for r_i in (0, 1, 2):  # widen Name / Designation / DOJ value cells so long names fit
        a, b2 = ann.cell(r_i, 2), ann.cell(r_i, 3)
        if a._tc is not b2._tc:
            m = a.merge(b2)
            for extra in m.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
    _cell_set(ann.cell(0, 2), emp.name)
    _cell_set(ann.cell(1, 2), emp.designation)
    _cell_set(ann.cell(2, 2), fmt_long(joining_date))
    # headers: template had them swapped; Yearly first, then Monthly
    _cell_set(ann.cell(4, 2), "Compensation (Yearly)")
    _cell_set(ann.cell(4, 3), "Compensation (Monthly)")
    row_map = {5: "basic", 6: "hra", 7: "additional", 8: "total_a", 10: "pf", 11: "gratuity",
               12: "pt", 13: "total_b", 14: "ctc"}
    for r_i, key in row_map.items():
        _cell_set(ann.cell(r_i, 2), inr(b["yearly"][key]))
        _cell_set(ann.cell(r_i, 3), inr(b["monthly"][key]))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF conversion (Word on Windows via docx2pdf, else LibreOffice)
# ---------------------------------------------------------------------------
def find_soffice() -> str | None:
    for c in ["soffice", "libreoffice",
              r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]:
        p = shutil.which(c) or (c if os.path.exists(c) else None)
        if p:
            return p
    return None


def docx_to_pdf_many(docx_paths: list[Path], out_dir: Path) -> list[Path]:
    """Convert several docx files to PDF. Returns list of created PDFs."""
    out_dir.mkdir(parents=True, exist_ok=True)
    if os.name == "nt":
        try:
            from docx2pdf import convert  # uses MS Word
            for p in docx_paths:
                convert(str(p), str(out_dir / (p.stem + ".pdf")))
            return [out_dir / (p.stem + ".pdf") for p in docx_paths]
        except Exception:
            pass
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("No PDF converter found. Install MS Word (with docx2pdf) or LibreOffice.")
    for i in range(0, len(docx_paths), 20):
        chunk = docx_paths[i:i + 20]
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir)]
                       + [str(p) for p in chunk], check=True, capture_output=True, timeout=600)
    return [out_dir / (p.stem + ".pdf") for p in docx_paths]


def docx_bytes_to_pdf(docx_bytes: bytes, stem: str) -> bytes:
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / f"{stem}.docx"
        src.write_bytes(docx_bytes)
        pdf = docx_to_pdf_many([src], Path(td) / "pdf")[0]
        return pdf.read_bytes()


# ---------------------------------------------------------------------------
# Offer register (keeps doc IDs unique across cycles)
# ---------------------------------------------------------------------------
REGISTER_COLS = ["Doc ID", "Year", "Seq", "Employee Name", "Designation", "Annual CTC", "Basic (Yearly)",
                 "Letter Date", "Joining Date", "File Name", "Generated On", "Status", "Remarks"]


def make_doc_id(year: int, seq: int, prefix: str = DOC_ID_PREFIX) -> str:
    return f"{prefix}/{year}/{seq}"


def safe_filename(doc_id: str, name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "-", f"{doc_id}_{name}").strip()


class Register:
    def __init__(self, path: str | os.PathLike):
        self.path = Path(path)
        if not self.path.exists():
            self.path.parent.mkdir(parents=True, exist_ok=True)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Register"
            ws.append(REGISTER_COLS)
            for c in ws[1]:
                c.font = openpyxl.styles.Font(bold=True)
            wb.save(self.path)

    def rows(self) -> list[dict]:
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Register"]
        hdr = [c.value for c in ws[1]]
        return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

    def next_seq(self, year: int, start: int = DOC_ID_START) -> int:
        seqs = [int(r["Seq"]) for r in self.rows() if r.get("Year") and int(r["Year"]) == year and r.get("Seq")]
        return max(seqs) + 1 if seqs else start

    def find(self, doc_id: str) -> dict | None:
        return next((r for r in self.rows() if r["Doc ID"] == doc_id), None)

    def add(self, entries: list[dict]):
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Register"]
        existing = {r[0].value: r[0].row for r in ws.iter_rows(min_row=2) if r[0].value}
        for e in entries:
            vals = [e.get(c) for c in REGISTER_COLS]
            if e["Doc ID"] in existing:  # re-issue: overwrite the row
                for ci, v in enumerate(vals, start=1):
                    ws.cell(existing[e["Doc ID"]], ci, v)
            else:
                ws.append(vals)
        widths = [24, 7, 6, 30, 32, 13, 14, 14, 14, 50, 18, 12, 40]
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        wb.save(self.path)


def register_entry(emp: Employee, doc_id: str, year: int, seq: int, letter_date: date, joining_date: date,
                   file_name: str, status: str = "Issued") -> dict:
    return {
        "Doc ID": doc_id, "Year": year, "Seq": seq, "Employee Name": emp.name, "Designation": emp.designation,
        "Annual CTC": emp.ctc, "Basic (Yearly)": emp.breakup["yearly"]["basic"],
        "Letter Date": letter_date, "Joining Date": joining_date, "File Name": file_name,
        "Generated On": datetime.now().strftime("%Y-%m-%d %H:%M"), "Status": status,
        "Remarks": "; ".join(emp.flags),
    }
